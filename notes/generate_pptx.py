import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation(output_path):
    prs = Presentation()
    
    # Use widescreen 16:9 aspect ratio
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Colors
    BG_COLOR = RGBColor(30, 30, 46)       # Dark slate blue
    TITLE_COLOR = RGBColor(166, 227, 161) # Pastel Green (Aya/Rust theme vibe)
    TEXT_COLOR = RGBColor(205, 214, 244)  # Light gray-white
    CYAN_ACCENT = RGBColor(137, 180, 250) # Pastel blue-cyan
    
    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        
    blank_layout = prs.slide_layouts[6]
    
    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "AI-FIREWALL"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR
    
    p2 = tf.add_paragraph()
    p2.text = "High-Performance, Self-Healing Network Security with Rust & eBPF"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Arial"
    p2.font.size = Pt(24)
    p2.font.color.rgb = CYAN_ACCENT
    
    p3 = tf.add_paragraph()
    p3.text = "\nTechnical Presentation"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Arial"
    p3.font.size = Pt(18)
    p3.font.color.rgb = TEXT_COLOR
    
    # Helper to add standard slide content
    def add_content_slide(title, bullets):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide)
        
        # Title Box
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.33), Inches(1.0))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(40)
        p_title.font.bold = True
        p_title.font.color.rgb = TITLE_COLOR
        
        # Content Box
        body_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(4.5))
        tf_body = body_box.text_frame
        tf_body.word_wrap = True
        
        for idx, bullet in enumerate(bullets):
            p = tf_body.paragraphs[0] if idx == 0 else tf_body.add_paragraph()
            p.text = bullet
            p.font.name = "Arial"
            p.font.size = Pt(22)
            p.font.color.rgb = TEXT_COLOR
            p.space_after = Pt(20)
            p.level = 0
            
    # -------------------------------------------------------------
    # SLIDE 2: The Performance & Security Dilemma
    # -------------------------------------------------------------
    add_content_slide(
        "Why eBPF and Rust?",
        [
            "• Traditional userspace firewalls require costly packet copies from kernel to userspace.",
            "• Traffic floods (DDoS) easily cause CPU starvation and high network latency.",
            "• eBPF runs sandboxed code directly inside the Linux kernel at near-hardware speeds.",
            "• Rust enforces memory safety in kernel-space (#![no_std]) and offers Tokio for userspace."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 3: Three-Tier Architecture
    # -------------------------------------------------------------
    add_content_slide(
        "Project Architecture",
        [
            "• ai-firewall-common: Shared telemetry structure definitions (e.g. PacketEvent).",
            "• ai-firewall-ebpf: The kernel-space code (no standard library, purely packet interception).",
            "• ai-firewall: The userspace controller (loads bytecode, manages maps, runs unblock scheduler).",
            "• Shared Maps: Enables bi-directional communication between kernel and userspace."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 4: Kernel Interception with XDP
    # -------------------------------------------------------------
    add_content_slide(
        "Safe Kernel Packet Interception",
        [
            "• XDP Hook intercepts packets directly at the network card driver level.",
            "• eBPF provides raw memory pointers: ctx.data() and ctx.data_end().",
            "• The Linux Kernel Verifier guarantees that memory access is safe and bounds-checked.",
            "• Network-types crate parses Ethernet & IPv4 headers cleanly without hand-written shifting."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 5: Shared State via eBPF Maps
    # -------------------------------------------------------------
    add_content_slide(
        "Shared State via eBPF Maps",
        [
            "• BLOCKLIST (HashMap): Key is source IP, value is status. Matches are dropped instantly.",
            "• EVENTS (RingBuf): A 16MB lockless queue to stream telemetry to userspace.",
            "• Zero-Copy Telemetry: Structured binary streaming avoids string-formatting overhead.",
            "• XDP_DROP action drops malicious packets before OS network stack allocation."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 6: Real-Time Threat Detection
    # -------------------------------------------------------------
    add_content_slide(
        "Self-Healing: Auto-Blocking Loop",
        [
            "• Userspace daemon reads telemetry events asynchronously from the Ring Buffer.",
            "• Maintains a real-time counter of packet frequencies per source IP address.",
            "• Auto-Block Rule: If an IP exceeds 100 packets, it is dynamically written into the BLOCKLIST.",
            "• Subsequent packets from the offending IP are dropped immediately in kernel-space."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 7: The TTL Sweeper
    # -------------------------------------------------------------
    add_content_slide(
        "The TTL Sweeper (Self-Healing)",
        [
            "• Blocked IPs must not remain blocked forever (prevents accidental permanent lockouts).",
            "• Userspace keeps an internal tracker mapping IP to the time of block (Instant).",
            "• A background TTL cleaner thread sweeps the blocked IPs every 5 seconds.",
            "• Expired IPs (> 60 seconds) are automatically deleted from the eBPF BLOCKLIST map."
        ]
    )
    
    # -------------------------------------------------------------
    # SLIDE 8: Summary & Next Steps
    # -------------------------------------------------------------
    add_content_slide(
        "Summary & Next Steps",
        [
            "• Microsecond Latency: Packet dropping happens before the kernel handles packet memory.",
            "• Zero CPU Starvation: Protects host applications during volumetric network attacks.",
            "• Next Steps: Support IPv6, parse ICMP headers, and add Deep Packet Inspection.",
            "• AI Rule Engine: Use lightweight ML models in userspace to adapt threshold dynamically."
        ]
    )
    
    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out = "/home/netrunner/Projects/ai-firewall/notes/ai_firewall_presentation.pptx"
    if len(sys.argv) > 1:
        out = sys.argv[1]
    create_presentation(out)
